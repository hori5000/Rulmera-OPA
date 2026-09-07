from pathlib import Path
from typing import Iterable
import csv
import re

import fitz
from docx import Document
from openpyxl import load_workbook

from .models import SourceBlock

TEXT_SUFFIXES = {".md", ".txt"}
SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".csv"}

def _base_meta(path: Path) -> dict:
    return {
        "source_file": path.name,
        "source_path": str(path.resolve()),
        "source_type": path.suffix.lower().lstrip("."),
    }

def parse_text(path: Path) -> Iterable[SourceBlock]:
    text = None
    for enc in ("utf-8", "utf-8-sig", "cp949"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise ValueError(f"Cannot decode text file: {path}")

    meta = _base_meta(path)

    if path.suffix.lower() != ".md":
        if text.strip():
            yield SourceBlock(text.strip(), meta)
        return

    heading = ""
    buf: list[str] = []

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m and buf:
            block = "\n".join(buf).strip()
            if block:
                yield SourceBlock(block, {**meta, "heading": heading})
            buf = []

        if m:
            heading = m.group(2).strip()
        buf.append(line)

    if buf:
        block = "\n".join(buf).strip()
        if block:
            yield SourceBlock(block, {**meta, "heading": heading})

def parse_pdf(path: Path) -> Iterable[SourceBlock]:
    doc = fitz.open(path)
    base = _base_meta(path)
    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if text:
            yield SourceBlock(text, {**base, "page": idx})

def parse_docx(path: Path) -> Iterable[SourceBlock]:
    doc = Document(path)
    base = _base_meta(path)
    heading = ""
    buf: list[str] = []

    def flush():
        nonlocal buf
        if not buf:
            return None
        text = "\n".join(buf).strip()
        buf = []
        if not text:
            return None
        return SourceBlock(text, {**base, "heading": heading})

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if p.style and p.style.name.lower().startswith("heading"):
            block = flush()
            if block:
                yield block
            heading = t
        buf.append(t)

    block = flush()
    if block:
        yield block

    for table_idx, table in enumerate(doc.tables, start=1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        text = "\n".join(rows).strip()
        if text:
            yield SourceBlock(
                text,
                {**base, "heading": heading, "table_index": table_idx},
            )

def parse_xlsx(path: Path, rows_per_block: int = 20) -> Iterable[SourceBlock]:
    wb = load_workbook(path, read_only=True, data_only=True)
    base = _base_meta(path)

    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue

        header_idx = None
        for i, row in enumerate(values):
            if any(v not in (None, "") for v in row):
                header_idx = i
                break

        if header_idx is None:
            continue

        header_raw = values[header_idx]
        headers = [
            str(v).strip() if v not in (None, "") else f"COL_{j+1}"
            for j, v in enumerate(header_raw)
        ]

        data = values[header_idx + 1:]

        for start in range(0, len(data), rows_per_block):
            rows = data[start:start + rows_per_block]
            lines = [f"Columns: {' | '.join(headers)}"]
            real_rows = []

            for offset, row in enumerate(rows):
                excel_row = header_idx + 2 + start + offset
                if not any(v not in (None, "") for v in row):
                    continue

                fields = []
                for j, value in enumerate(row):
                    if value in (None, ""):
                        continue
                    key = headers[j] if j < len(headers) else f"COL_{j+1}"
                    fields.append(f"{key}={value}")

                lines.append(f"Row {excel_row}: " + " | ".join(fields))
                real_rows.append(excel_row)

            if real_rows:
                yield SourceBlock(
                    "\n".join(lines),
                    {
                        **base,
                        "sheet": ws.title,
                        "row_start": min(real_rows),
                        "row_end": max(real_rows),
                        "header_row": header_idx + 1,
                    },
                )

def parse_csv_file(path: Path, rows_per_block: int = 30) -> Iterable[SourceBlock]:
    base = _base_meta(path)
    text = None

    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        return

    rows = list(csv.reader(text.splitlines()))
    if not rows:
        return

    headers = rows[0]

    for start in range(1, len(rows), rows_per_block):
        block = rows[start:start + rows_per_block]
        lines = [f"Columns: {' | '.join(headers)}"]

        for row_no, row in enumerate(block, start=start + 1):
            fields = []
            for j, value in enumerate(row):
                if value == "":
                    continue
                key = headers[j] if j < len(headers) else f"COL_{j+1}"
                fields.append(f"{key}={value}")
            lines.append(f"Row {row_no}: " + " | ".join(fields))

        yield SourceBlock(
            "\n".join(lines),
            {
                **base,
                "row_start": start + 1,
                "row_end": start + len(block),
            },
        )

def parse_file(path: Path) -> list[SourceBlock]:
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED:
        return []
    if suffix in TEXT_SUFFIXES:
        return list(parse_text(path))
    if suffix == ".pdf":
        return list(parse_pdf(path))
    if suffix == ".docx":
        return list(parse_docx(path))
    if suffix == ".xlsx":
        return list(parse_xlsx(path))
    if suffix == ".csv":
        return list(parse_csv_file(path))

    return []
